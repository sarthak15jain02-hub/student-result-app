from flask import Flask, render_template, request, redirect, url_for, session, send_file
from fpdf import FPDF
import openpyxl, io, os

app = Flask(__name__)
app.secret_key = 'devops2024secret'

# Admin credentials
ADMIN = {'username': 'admin', 'password': 'admin123'}

# Admin profile
PROFILE = {
    'name': 'Prof. Harshali Vihire',
    'email': 'admin@school.edu',
    'role': 'Administrator',
    'department': 'Computer Science',
    'phone': '+91 98765 43210'
}

# In-memory student store
students = []

def get_grade(marks):
    if marks >= 80: return 'A'
    elif marks >= 60: return 'B'
    elif marks >= 40: return 'C'
    else: return 'F'

def login_required(f):
    from functools import wraps
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'logged_in' not in session:
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated

@app.route('/', methods=['GET', 'POST'])
def login():
    error = None
    if request.method == 'POST':
        if request.form['username'] == ADMIN['username'] and request.form['password'] == ADMIN['password']:
            session['logged_in'] = True
            return redirect(url_for('dashboard'))
        error = 'Invalid username or password'
    return render_template('login.html', error=error)

@app.route('/dashboard')
@login_required
def dashboard():
    total = len(students)
    passed = sum(1 for s in students if s['marks'] >= 40)
    failed = total - passed
    grades = {'A': 0, 'B': 0, 'C': 0, 'F': 0}
    for s in students:
        grades[s['grade']] += 1
    return render_template('dashboard.html', total=total, passed=passed, failed=failed, grades=grades, students=students)

@app.route('/add', methods=['POST'])
@login_required
def add_student():
    name = request.form['name'].strip()
    subject = request.form['subject'].strip()
    marks = int(request.form['marks'])
    grade = get_grade(marks)
    students.append({'id': len(students)+1, 'name': name, 'subject': subject, 'marks': marks, 'grade': grade})
    return redirect(url_for('dashboard'))

@app.route('/delete/<int:sid>')
@login_required
def delete_student(sid):
    global students
    students = [s for s in students if s['id'] != sid]
    return redirect(url_for('dashboard'))

@app.route('/profile')
@login_required
def profile():
    return render_template('profile.html', profile=PROFILE)

@app.route('/export/pdf')
@login_required
def export_pdf():
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font('Helvetica', 'B', 16)
    pdf.cell(0, 12, 'Student Result Report', ln=True, align='C')
    pdf.set_font('Helvetica', 'B', 11)
    pdf.set_fill_color(66, 133, 244)
    pdf.set_text_color(255, 255, 255)
    for col, w in [('ID',15),('Name',55),('Subject',55),('Marks',30),('Grade',25)]:
        pdf.cell(w, 10, col, border=1, fill=True)
    pdf.ln()
    pdf.set_font('Helvetica', '', 10)
    pdf.set_text_color(0, 0, 0)
    for s in students:
        pdf.set_fill_color(240, 240, 240) if s['id'] % 2 == 0 else pdf.set_fill_color(255, 255, 255)
        pdf.cell(15, 9, str(s['id']), border=1, fill=True)
        pdf.cell(55, 9, s['name'], border=1, fill=True)
        pdf.cell(55, 9, s['subject'], border=1, fill=True)
        pdf.cell(30, 9, str(s['marks']), border=1, fill=True)
        pdf.cell(25, 9, s['grade'], border=1, fill=True)
        pdf.ln()
    buf = io.BytesIO(pdf.output())
    return send_file(buf, mimetype='application/pdf', download_name='results.pdf', as_attachment=True)

@app.route('/export/excel')
@login_required
def export_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Results'
    ws.append(['ID', 'Name', 'Subject', 'Marks', 'Grade'])
    for s in students:
        ws.append([s['id'], s['name'], s['subject'], s['marks'], s['grade']])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', download_name='results.xlsx', as_attachment=True)

@app.route('/health')
def health():
    return {'status': 'UP', 'students': len(students)}, 200

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)